"""Unit tests for document loading and Graphiti document ingestion.

Loader tests use tmp_path fixtures so they stay correct even if the real
mock content under documents/ changes wording. A separate end-to-end test
loads the real documents/ directory to confirm every mock file (sql, py,
java, txt, docx, pptx, image) is actually parseable.
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import AsyncMock

import pytest
from docx import Document as DocxDocument
from graphiti_core.nodes import EpisodeType
from PIL import Image
from pptx import Presentation

from src.config import get_settings
from src.document_loader import (
    DocumentRecord,
    UnsupportedDocumentTypeError,
    load_document,
    load_documents_dir,
)
from src.graphiti_engine import DOCUMENT_GROUP_ID, GraphitiEngine

REPO_DOCUMENTS_DIR = Path(__file__).resolve().parent.parent / "documents"


@pytest.fixture
def engine(mocker) -> GraphitiEngine:
    mock_graphiti_cls = mocker.patch("src.graphiti_engine.Graphiti")
    mock_instance = mock_graphiti_cls.return_value
    mock_instance.add_episode = AsyncMock()
    mock_instance.close = AsyncMock()
    return GraphitiEngine(settings=get_settings())


class TestTextLoaders:
    @pytest.mark.parametrize(
        ("filename", "expected_type"),
        [
            ("notes.txt", "text"),
            ("schema.sql", "sql"),
            ("script.py", "python"),
            ("Listener.java", "java"),
        ],
    )
    def test_reads_plain_text_files_by_extension(self, tmp_path, filename, expected_type):
        file_path = tmp_path / filename
        file_path.write_text("hello project", encoding="utf-8")

        record = load_document(file_path, tmp_path)

        assert record.doc_type == expected_type
        assert record.content_text == "hello project"
        assert record.doc_id == filename

    def test_unsupported_extension_raises(self, tmp_path):
        file_path = tmp_path / "archive.zip"
        file_path.write_bytes(b"not really a zip")

        with pytest.raises(UnsupportedDocumentTypeError):
            load_document(file_path, tmp_path)


class TestDocxLoader:
    def test_extracts_paragraph_text(self, tmp_path):
        doc = DocxDocument()
        doc.add_heading("Project Charter", level=1)
        doc.add_paragraph("Objective: build a temporal knowledge graph.")
        docx_path = tmp_path / "charter.docx"
        doc.save(docx_path)

        record = load_document(docx_path, tmp_path)

        assert record.doc_type == "docx"
        assert "Project Charter" in record.content_text
        assert "temporal knowledge graph" in record.content_text


class TestPptxLoader:
    def test_extracts_slide_text(self, tmp_path):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Architecture"
        slide.placeholders[1].text = "Gmail -> Graphiti -> Neo4j"
        pptx_path = tmp_path / "overview.pptx"
        prs.save(pptx_path)

        record = load_document(pptx_path, tmp_path)

        assert record.doc_type == "pptx"
        assert "Architecture" in record.content_text
        assert "Neo4j" in record.content_text


class TestImageLoader:
    def test_uses_caption_sidecar_and_dimensions(self, tmp_path):
        image_path = tmp_path / "diagram.png"
        Image.new("RGB", (200, 100), "white").save(image_path)
        (tmp_path / "diagram.caption.txt").write_text(
            "A diagram of the ingestion pipeline.", encoding="utf-8"
        )

        record = load_document(image_path, tmp_path)

        assert record.doc_type == "image"
        assert "200x100" in record.content_text
        assert "A diagram of the ingestion pipeline." in record.content_text

    def test_missing_caption_is_handled_gracefully(self, tmp_path):
        image_path = tmp_path / "no_caption.png"
        Image.new("RGB", (50, 50), "white").save(image_path)

        record = load_document(image_path, tmp_path)

        assert "no caption available" in record.content_text


class TestLoadDocumentsDir:
    def test_skips_caption_sidecars_and_hidden_files(self, tmp_path):
        (tmp_path / "notes.txt").write_text("hi", encoding="utf-8")
        Image.new("RGB", (10, 10)).save(tmp_path / "pic.png")
        (tmp_path / "pic.caption.txt").write_text("a picture", encoding="utf-8")
        (tmp_path / ".hidden").write_text("secret", encoding="utf-8")

        records = load_documents_dir(tmp_path)
        doc_ids = {r.doc_id for r in records}

        assert doc_ids == {"notes.txt", "pic.png"}

    def test_real_project_documents_all_load(self):
        """End-to-end check that every mock file under documents/ is parseable."""
        records = load_documents_dir(REPO_DOCUMENTS_DIR)
        doc_types = {r.doc_type for r in records}

        assert doc_types == {"sql", "python", "java", "text", "docx", "pptx", "image"}
        assert all(r.content_text.strip() for r in records)


class TestDocumentIngestion:
    @pytest.mark.asyncio
    async def test_ingest_document_uses_shared_document_group_and_text_type(self, engine):
        doc = DocumentRecord(
            doc_id="schema/neo4j_schema.sql",
            filename="neo4j_schema.sql",
            doc_type="sql",
            content_text="CREATE TABLE emails (...);",
            modified_at=datetime.now(timezone.utc),
        )

        await engine.ingest_document(doc)

        engine.graphiti.add_episode.assert_awaited_once()
        _, kwargs = engine.graphiti.add_episode.await_args

        assert kwargs["name"] == "Document: neo4j_schema.sql (sql)"
        assert "CREATE TABLE emails" in kwargs["episode_body"]
        assert kwargs["source"] == EpisodeType.text
        assert kwargs["source_description"] == "document:sql"
        assert kwargs["group_id"] == DOCUMENT_GROUP_ID
